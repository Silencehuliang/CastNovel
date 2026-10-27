"""导出:TXT / Markdown / DOCX(python-docx 可选)。"""
import re


def _safe_name(title):
    return re.sub(r'[\\/:*?"<>|]+', "_", title or "novel").strip() or "novel"


def export_text(project, chapters, fmt="md"):
    title = project.get("title") or "未命名作品"
    author_line = project.get("genre", "")
    if fmt == "txt":
        parts = [title, f"类型:{author_line}" if author_line else "", ""]
        for ch in chapters:
            parts.append(f"第{ch['idx']}章 {ch.get('title', '')}")
            parts.append(ch.get("content") or "(未写)")
            parts.append("")
        return "\n".join(parts), f"{_safe_name(title)}.txt", "text/plain; charset=utf-8"

    parts = [f"# {title}", ""]
    if author_line:
        parts += [f"> 类型:{author_line}", ""]
    if project.get("premise"):
        parts += [f"**梗概:**{project['premise']}", ""]
    for ch in chapters:
        parts.append(f"## 第{ch['idx']}章 {ch.get('title', '')}")
        parts.append("")
        parts.append(ch.get("content") or "(未写)")
        parts.append("")
    return "\n".join(parts), f"{_safe_name(title)}.md", "text/markdown; charset=utf-8"


def export_docx(project, chapters):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("导出 DOCX 需要安装 python-docx:pip install python-docx")
    doc = Document()
    title = project.get("title") or "未命名作品"
    doc.add_heading(title, level=0)
    if project.get("premise"):
        doc.add_paragraph(project["premise"])
    for ch in chapters:
        doc.add_heading(f"第{ch['idx']}章 {ch.get('title', '')}", level=1)
        for para in (ch.get("content") or "(未写)").splitlines():
            if para.strip():
                doc.add_paragraph(para.strip())
    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"{_safe_name(title)}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def export_epub(project, chapters):
    """最小可用的 EPUB 3(零依赖,zipfile 实现):mimetype + container + nav + opf + 章节 xhtml。"""
    import io
    import uuid
    import zipfile
    from xml.sax.saxutils import escape

    title = project.get("title") or "未命名作品"
    book = [ch for ch in chapters if (ch.get("content") or "").strip()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?>'
            '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        manifest, spine, toc = [], [], []
        for ch in book:
            fname = f"ch{ch['idx']}.xhtml"
            heading = escape(f"第{ch['idx']}章 {ch.get('title', '')}")
            paras = "".join(
                f"<p>{escape(p.strip())}</p>"
                for p in (ch.get("content") or "").splitlines() if p.strip()
            ) or "<p>(空)</p>"
            z.writestr(
                f"OEBPS/{fname}",
                '<?xml version="1.0" encoding="utf-8"?>'
                '<!DOCTYPE html>'
                '<html xmlns="http://www.w3.org/1999/xhtml">'
                f'<head><title>{heading}</title><meta charset="utf-8"/></head>'
                f"<body><h2>{heading}</h2>{paras}</body></html>",
            )
            manifest.append(f'<item id="c{ch["idx"]}" href="{fname}" media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="c{ch["idx"]}"/>')
            toc.append(f'<li><a href="{fname}">{heading}</a></li>')
        z.writestr(
            "OEBPS/nav.xhtml",
            '<?xml version="1.0" encoding="utf-8"?>'
            '<!DOCTYPE html>'
            '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
            '<head><title>目录</title><meta charset="utf-8"/></head>'
            f'<body><nav epub:type="toc"><h1>{escape(title)}</h1><ol>{"".join(toc)}</ol></nav></body></html>',
        )
        z.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0" encoding="utf-8"?>'
            '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            f'<dc:identifier id="bookid">urn:uuid:{uuid.uuid4()}</dc:identifier>'
            f"<dc:title>{escape(title)}</dc:title>"
            "<dc:language>zh</dc:language>"
            "<meta property=\"dcterms:modified\">2026-01-01T00:00:00Z</meta>"
            "</metadata>"
            '<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
            f'{"".join(manifest)}</manifest>'
            f'<spine>{"".join(spine)}</spine>'
            "</package>",
        )
    return buf.getvalue(), f"{_safe_name(title)}.epub", "application/epub+zip"
